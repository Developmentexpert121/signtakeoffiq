#!/usr/bin/env python3
"""
xlsx-to-json.py — convert a Takeoff XLSX into scorer-ready {rooms, signs} JSON.

The app's XLSX export ("Tower_Dist_Takeoff_<date>.xlsx") and a human-produced
"golden copy" saved in that same layout share one sheet — **Takeoff** — with the
columns:

  Floor | Room # | Room Name | Sign Type | Size | Qty | ADA | Notes | ... | Extended

This converts that sheet into the shape score-takeoff.mjs eats:

  { "rooms": [{roomNumber, roomName, level}], "signs": [{roomNumber, signType, qty, level}] }

So the SAME tool turns both the AI export AND the hand-built golden copy into
comparable JSON — feed the two JSON files to score-takeoff.mjs and you get a number.

Usage:
  python3 scripts/takeoff-eval/xlsx-to-json.py INPUT.xlsx [-o OUTPUT.json] [--sheet Takeoff]

Requires: openpyxl  (pip install openpyxl)
"""
import argparse
import json
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

# Header label -> our field. Matched case-insensitively, trimmed.
HEADER_ALIASES = {
    "floor": "level",
    "level": "level",
    "room #": "roomNumber",
    "room#": "roomNumber",
    "room number": "roomNumber",
    "room name": "roomName",
    "sign type": "signType",
    "type": "signType",
    "qty": "qty",
    "quantity": "qty",
}

# Rows whose Room # / Sign Type look like these are structure, not data.
SKIP_TOKENS = {"", "subtotal", "total", "grand total", "installation"}


def norm(v):
    return "" if v is None else str(v).strip()


def find_header_row(ws, max_scan=10):
    """Find the row that contains 'Sign Type' + 'Qty' — that's the real header."""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        labels = [norm(c).lower() for c in next(ws.iter_rows(min_row=r, max_row=r, values_only=True))]
        if "sign type" in labels and ("qty" in labels or "quantity" in labels):
            return r, labels
    return None, None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="path to the .xlsx takeoff")
    ap.add_argument("-o", "--output", help="output .json (default: stdout)")
    ap.add_argument("--sheet", default=None, help="sheet name (default: first sheet named 'Takeoff', else first sheet)")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.input, data_only=True)
    sheet_name = args.sheet or ("Takeoff" if "Takeoff" in wb.sheetnames else wb.sheetnames[0])
    if sheet_name not in wb.sheetnames:
        sys.exit(f"sheet {sheet_name!r} not found. available: {wb.sheetnames}")
    ws = wb[sheet_name]

    header_row, labels = find_header_row(ws)
    if header_row is None:
        sys.exit("could not locate a header row containing 'Sign Type' and 'Qty'")

    # Map column index -> our field name.
    col_field = {}
    for idx, lab in enumerate(labels):
        field = HEADER_ALIASES.get(lab)
        if field:
            col_field[idx] = field

    if "signType" not in col_field.values():
        sys.exit("no 'Sign Type' column found")

    rooms = {}   # roomNumber -> {roomNumber, roomName, level}
    signs = []
    skipped = 0

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rec = {}
        for idx, field in col_field.items():
            if idx < len(row):
                rec[field] = norm(row[idx])

        room_no = rec.get("roomNumber", "")
        sign_type = rec.get("signType", "")

        # Skip section headers, subtotals, blank rows.
        if room_no.lower() in SKIP_TOKENS or sign_type.lower() in SKIP_TOKENS:
            skipped += 1
            continue
        if not sign_type:
            skipped += 1
            continue

        level = rec.get("level", "")
        room_name = rec.get("roomName", "")

        if room_no and room_no not in rooms:
            rooms[room_no] = {"roomNumber": room_no, "roomName": room_name, "level": level}

        qty_raw = rec.get("qty", "")
        try:
            qty = int(float(qty_raw)) if qty_raw not in ("", "—", "-") else 1
        except (ValueError, TypeError):
            qty = 1

        signs.append({
            "roomNumber": room_no,
            "roomName": room_name,
            "signType": sign_type,
            "qty": qty,
            "level": level,
        })

    out = {
        "name": f"{args.input} :: {sheet_name}",
        "rooms": list(rooms.values()),
        "signs": signs,
    }

    text = json.dumps(out, indent=2)
    if args.output:
        with open(args.output, "w") as f:
            f.write(text + "\n")
        total_qty = sum(s["qty"] for s in signs)
        print(f"wrote {args.output}: {len(rooms)} rooms, {len(signs)} sign rows "
              f"({total_qty} total qty), {skipped} non-data rows skipped", file=sys.stderr)
    else:
        print(text)


if __name__ == "__main__":
    main()
